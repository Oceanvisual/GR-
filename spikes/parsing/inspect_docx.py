"""Spike T072: what a real regulation DOCX looks like to the parser.

Usage: uv run python -I spikes/parsing/inspect_docx.py <file.docx>
"""

import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

W_T = qn("w:t")
W_DEL = qn("w:del")
W_INS = qn("w:ins")

NOISE_PATTERNS = {
    "consultant_banner": r"КонсультантПлюс",
    "consultant_url": r"www\.consultant\.ru",
    "save_date": r"Дата сохранения",
    "change_list": r"Список изменяющих документов",
    "in_edition": r"\(в ред\.",
    "introduced": r"\((?:п\.|абзац|подп?\.)?.{0,20}введ[её]н",
    "repealed": r"[Уу]тратил[аио]? силу",
}
NUMBERING = {
    "roman": re.compile(r"^(?P<n>[IVXL]+)\.\s"),
    "decimal": re.compile(r"^(?P<n>\d+(?:\.\d+)*)\.?\s"),
    "paren_digit": re.compile(r"^(?P<n>\d+)\)\s"),
    "paren_letter": re.compile(r"^(?P<n>[а-я])\)\s"),
    "dash": re.compile(r"^[-–—]\s"),
}


def xml_text(paragraph_element: etree._Element) -> str:
    """All w:t in the paragraph except deleted runs (includes w:ins, hyperlinks, fields)."""
    parts = []
    for node in paragraph_element.iter(W_T):
        if any(ancestor.tag == W_DEL for ancestor in node.iterancestors()):
            continue
        parts.append(node.text or "")
    return "".join(parts)


def classify(line: str) -> tuple[str, str] | None:
    for name, pattern in NUMBERING.items():
        if m := pattern.match(line):
            return name, m.groupdict().get("n", "")
    return None


def plausible_next(prev: tuple[int, ...] | None, cur: tuple[int, ...]) -> bool:
    if prev is None:
        return True
    for depth in range(len(cur)):
        if cur[:depth] == prev[:depth] and len(cur) == depth + 1:
            base = prev[depth] if depth < len(prev) else 0
            if cur[depth] == base + 1:
                return True
    return False


def main(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        header_footer = {
            name: " ".join(t.text or "" for t in etree.fromstring(archive.read(name)).iter(W_T))
            for name in names
            if re.match(r"word/(header|footer)\d+\.xml", name)
        }
    doc = Document(str(path))
    body = doc.element.body
    paragraphs = [p for p in body.iter(qn("w:p"))]
    in_tables = sum(1 for p in paragraphs if any(a.tag == qn("w:tbl") for a in p.iterancestors()))
    numpr = sum(1 for p in paragraphs if p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None)
    ins = sum(1 for _ in body.iter(W_INS))
    dels = sum(1 for _ in body.iter(W_DEL))
    hyperlinks = sum(1 for _ in body.iter(qn("w:hyperlink")))
    fields = sum(1 for _ in body.iter(qn("w:fldSimple"))) + sum(
        1 for _ in body.iter(qn("w:instrText"))
    )

    api_texts = [p.text for p in doc.paragraphs]
    top_level = [p for p in body.iterchildren(qn("w:p"))]
    xml_texts = [xml_text(p) for p in top_level]
    diffs = [(a, b) for a, b in zip(api_texts, xml_texts, strict=True) if a != b]

    lines = [t.strip() for t in xml_texts if t.strip()]
    total_chars = sum(len(t) for t in lines)
    noise = {k: sum(1 for t in lines if re.search(p, t)) for k, p in NOISE_PATTERNS.items()}
    kinds = Counter(c[0] for t in lines if (c := classify(t)))

    decimal_seq = []
    prev: tuple[int, ...] | None = None
    implausible = []
    for t in lines:
        c = classify(t)
        if c and c[0] == "decimal":
            cur = tuple(int(x) for x in c[1].split("."))
            if not plausible_next(prev, cur):
                implausible.append((c[1], t[:70]))
            decimal_seq.append(cur)
            prev = cur
    dup = [n for n, k in Counter(decimal_seq).items() if k > 1]

    print(f"file: {path.name}")
    print(f"paragraphs: {len(paragraphs)} (top-level {len(top_level)}, in tables {in_tables})")
    print(f"tables: {len(doc.tables)}; numPr paragraphs: {numpr}; w:ins {ins}; w:del {dels}")
    print(f"hyperlinks: {hyperlinks}; fields: {fields}; chars: {total_chars}")
    print(f"header/footer text: {header_footer}")
    print(f"paragraph.text != xml walk: {len(diffs)}")
    for a, b in diffs[:5]:
        print(f"  api: {a[:90]!r}\n  xml: {b[:90]!r}")
    print(f"noise lines: {noise}")
    print(f"numbering kinds: {dict(kinds)}")
    print(
        f"decimal points: {len(decimal_seq)}; "
        f"duplicates: {['.'.join(map(str, d)) for d in dup][:15]}"
    )
    print(f"implausible continuations: {len(implausible)}")
    for n, t in implausible[:25]:
        print(f"  {n:>8} | {t}")
    print("--- first 25 non-empty lines:")
    for t in lines[:25]:
        print(f"  {t[:110]}")
    print("--- lines with noise (first 15):")
    for t in [t for t in lines if any(re.search(p, t) for p in NOISE_PATTERNS.values())][:15]:
        print(f"  {t[:140]}")


if __name__ == "__main__":
    main(Path(sys.argv[1]))
