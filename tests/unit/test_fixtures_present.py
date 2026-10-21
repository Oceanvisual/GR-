"""Synthetic fixtures exist and have the properties tests rely on (see make_fixtures.py)."""

import zipfile
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from docx import Document

SYNTHETIC = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic"

TEXT_FILES = [
    "regulation.txt",
    "injection.txt",
    "no_numbering.txt",
    "law_like.txt",
    "dup_numbering.txt",
    "scan.expected.txt",
]
DOCX_FILES = [
    "regulation_manual.docx",
    "regulation_autonum.docx",
    "multi_approved.docx",
    "spps_export.docx",
    "table.docx",
    "tracked_changes.docx",
    "styled_headings.docx",
]


def _pdf_text(name: str) -> list[str]:
    document = pdfium.PdfDocument(str(SYNTHETIC / name))
    try:
        return [page.get_textpage().get_text_range() for page in document]
    finally:
        document.close()


@pytest.mark.parametrize("name", TEXT_FILES)
def test_utf8_text_fixtures(name: str) -> None:
    assert (SYNTHETIC / name).read_text(encoding="utf-8").strip()


def test_cp1251_fixture_is_not_utf8() -> None:
    raw = (SYNTHETIC / "regulation_cp1251.txt").read_bytes()
    with pytest.raises(UnicodeDecodeError):
        raw.decode("utf-8")
    assert raw.decode("cp1251") == (SYNTHETIC / "regulation.txt").read_text(encoding="utf-8")


@pytest.mark.parametrize("name", DOCX_FILES)
def test_docx_fixtures_open(name: str) -> None:
    path = SYNTHETIC / name
    with zipfile.ZipFile(path) as archive:
        assert "word/document.xml" in archive.namelist()
    assert Document(str(path)).paragraphs


def test_autonum_docx_has_numbering_without_number_text() -> None:
    paragraphs = Document(str(SYNTHETIC / "regulation_autonum.docx")).paragraphs
    numbered = [p for p in paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None]
    assert len(numbered) >= 20
    assert all(not p.text[:1].isdigit() for p in numbered)


def test_text_pdf_has_text_layer() -> None:
    pages = _pdf_text("regulation.pdf")
    assert len(pages) >= 2
    assert all(len(page.strip()) > 500 for page in pages)


def test_headers_pdf_layout_traps() -> None:
    pages = _pdf_text("headers.pdf")
    assert len(pages) >= 3
    assert all(page.splitlines()[0].startswith("Управление охраны") for page in pages)
    assert all(page.rstrip().splitlines()[-1].startswith("Страница") for page in pages)
    lines = [line for page in pages for line in page.splitlines()]
    assert any(line.startswith("10 календарных дней") for line in lines)


def test_scan_pdf_has_no_text_layer() -> None:
    assert [page.strip() for page in _pdf_text("scan.pdf")] == [""]


def test_corrupted_pdf_is_truncated() -> None:
    raw = (SYNTHETIC / "corrupted.pdf").read_bytes()
    assert raw.startswith(b"%PDF-")
    assert b"%%EOF" not in raw


def test_html_declares_cp1251() -> None:
    raw = (SYNTHETIC / "regulation.html").read_bytes()
    assert b"charset=windows-1251" in raw
    assert "Общие положения" in raw.decode("cp1251")
